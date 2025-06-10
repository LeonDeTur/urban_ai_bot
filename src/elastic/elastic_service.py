import io

import pandas as pd
from tqdm import tqdm

from docx import Document
from elastic_transport import ObjectApiResponse
from elasticsearch import Elasticsearch
from elasticsearch.helpers import bulk
from fastapi import HTTPException
from loguru import logger

from src.dependencies import http_exception
from src.common.config.config import Config
from src.llm.llm_service import LlmService
from src.vectorizer.vectorizer_service import VectorizerService
from .doc_parser import doc_parser


class ElasticService:
    def __init__(self, config: Config, vectorizer_service: VectorizerService, llm_service: LlmService):
        self.client = Elasticsearch(hosts=[f"http://{config.get('ELASTIC_HOST')}:{config.get('ELASTIC_PORT')}"])
        self.config = config
        self.vectorizer_service = vectorizer_service
        self.llm_service = llm_service

    async def get_available_indexes(self) -> list[str]:

        all_indices = self.client.indices.get_alias(index="*")
        return [index for index in all_indices if not index.startswith('.') and not index.startswith('_')]

    async def create_index(self, index_name: str):

        try:
            resp = self.client.indices.create(index=index_name, body={
                "mappings": {
                    "properties": {
                        "body_vector": {
                            "type": "dense_vector",
                            "dims": 4096,
                            "index": True,
                            "similarity": "cosine"
                        },
                        "body": {
                            "type": "text"
                        },
                        "num_id": {
                            "type": "long"
                        }
                    }
                }})

            return resp.raw

        except Exception as e:
            raise http_exception(
                500,
                "Failed to create index",
                _input={"index_name": index_name},
                _detail={"error": e.__str__()}
            )

    async def delete_index(self, index_name: str):

        resp = self.client.options(ignore_status=[400, 404]).indices.delete(index=index_name)
        return resp.raw

    async def delete_documents_from_index(self, index_name: str) -> str:
        try:
            self.client.delete_by_query(index=index_name, body={"query": {"match_all": {}}})
            return f"Successfully deleted all documents from index {index_name}"
        except Exception as e:
            logger.exception(e)
            raise HTTPException(status_code=500, detail=e.__str__())

    async def search(self, embedding: list, index_name: str) -> ObjectApiResponse:

        query_body = {
            "knn": {
                "field": "body_vector",
                "query_vector": embedding,
                "k": 20,
                "num_candidates": 20
            },
            "_source": ["body"],
        }
        return self.client.search(index=index_name, body=query_body)

    async def get_last_index(self, index_name: str) -> int:
        query_body = {
            "size": 1,
            "sort": [
                {
                    "num_id": {
                        "order": "desc"
                    }
                }
            ]
        }
        try:
            last_id_data = self.client.search(index=index_name, body=query_body)
        except Exception as e:
            logger.exception(e)
            raise HTTPException(status_code=500, detail=e.__str__())
        if last_id_data.body["hits"]["hits"]:
            return last_id_data.body["hits"]["hits"][0]["_source"]["num_id"]
        return 0

    async def create_doc_to_upload(self, text: str, doc_id: int) -> dict[str, str, ]:

        vector = self.encode(text)
        return {
            "_id": str(doc_id),
            "num_id": doc_id,
            "body": text,
            "body_vector": vector,
        }

    async def create_table_to_upload(
            self,
            table_with_context: tuple[str , str, str],
            num_questions: int,
            last_doc_id: int
    ) -> tuple[list[dict[str, str | int]], int]:

        docs_to_add = []
        table_questions = await self.llm_service.generate_table_description(table_with_context[1], num_questions)
        text = "\n".join(table_with_context)
        for i in range(1, len(table_questions) + 2):
            if i > len(table_questions):
                docs_to_add.append(
                    {
                        "_id": str(last_doc_id + i),
                        "num_id": last_doc_id + i,
                        "body": text,
                        "body_vector": self.encode(text),
                    }
                )
            else:
                docs_to_add.append(
                    {
                        "_id": str(last_doc_id + i),
                        "num_id": last_doc_id + i,
                        "body": text,
                        "body_vector": self.encode(table_questions[i - 1]),
                    }
                )

        return docs_to_add, last_doc_id + len(table_questions) + 1

    async def upload_to_index(self, file: bytes, index_name: str, table_context_size: int, table_questions_num: int):
        if not self.client.indices.exists(index=index_name):
            await self.create_index(index_name)
        documents = []
        full_doc = Document(io.BytesIO(file))
        last_id = await self.get_last_index(index_name)
        logger.info(f"Started uploading documents to index {index_name} from id {last_id}")
        dock_blocks = [i for i in doc_parser.iter_contexts_for_vectorization(full_doc)]
        for index, entity in tqdm(enumerate(dock_blocks), total=len(dock_blocks), desc="Processing texts"):
            if entity[1] == "text":
                text = entity[0]
                doc_to_upload = await self.create_doc_to_upload(text, last_id)
                documents.append(doc_to_upload)
                last_id += 1
            elif entity[1] == "table":
                table_data = entity[0]
                try:
                    table_with_context = (
                        "\n".join([i[0] for i in dock_blocks[index - table_context_size: index]]),
                        table_data,
                        "\n".join([i[0] for i in dock_blocks[index: index - table_context_size]])
                    )
                except Exception as e:
                    logger.exception(e)
                    raise HTTPException(status_code=500, detail=e.__str__())
                docs, last_id = await self.create_table_to_upload(table_with_context, table_questions_num, last_id)
                documents += docs

        for table in tqdm(full_doc.tables, total=len(full_doc.tables) ,desc="Processint tables"):
            last_index = 0
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            table_df = pd.DataFrame(table_data[1:], columns=table_data[0])
            for index, row in table_df.iterrows():
                # table_data.append(row_data)
                # table_description = await self.llm_service.generate_table_description(table_data)
                vector = self.encode(str(row))
                doc = {
                    "_id": str(ids + index + 1),
                    "num_id": index,
                    "body": str(row),
                    "body_vector": vector,
                }
                documents.append(doc)
                last_index += 1
            ids += last_index

        if documents:
            bulk(self.client, documents, index=index_name)
        print("Finished")
        return index_name

    def encode(self, document: str) -> list:
        try:
            return self.vectorizer_service.embed(document)
        except Exception as e:
            raise HTTPException(500, str(e))
